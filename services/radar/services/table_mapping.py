"""Deterministic table snapshots and typed facts from mapped rows."""

from __future__ import annotations

import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
from typing import Any

from pydantic import BaseModel, Field, model_validator

from services.radar.domain.models import Fact
from services.radar.services.source_versions import resolve_quote


class TableMappingError(ValueError):
    pass


class TypedColumn(BaseModel):
    header: str = Field(min_length=1)
    predicate: str = Field(min_length=1)
    type: str = Field(pattern="^(date|decimal|enum)$")
    unit: str | None = None
    enum_values: list[str] | None = None

    @model_validator(mode="after")
    def complete(self):
        if self.type == "decimal" and not self.unit:
            raise ValueError("A decimal column requires a unit.")
        if self.type == "enum" and not self.enum_values:
            raise ValueError("An enum column requires allowed values.")
        return self


class TableMapping(BaseModel):
    sheet_name: str | None = None
    a1_range: str | None = None
    header_row: int = Field(default=1, ge=1)
    row_business_key: str = Field(min_length=1)
    columns: list[TypedColumn] = Field(min_length=1, max_length=20)

    @model_validator(mode="after")
    def unique_headers(self):
        names = [self.row_business_key, *(column.header for column in self.columns)]
        if len(names) != len(set(names)):
            raise ValueError("The row key and typed columns must have distinct headers.")
        return self


def csv_rows(data: bytes) -> list[list[Any]]:
    try:
        return list(csv.reader(StringIO(data.decode("utf-8-sig"), newline="")))
    except (UnicodeDecodeError, csv.Error) as exc:
        raise TableMappingError("CSV is not valid UTF-8 table data.") from exc


def xlsx_rows(data: bytes, sheet_name: str | None) -> tuple[list[list[Any]], list[list[Any]], str]:
    from openpyxl import load_workbook
    try:
        values_book = load_workbook(BytesIO(data), read_only=True, data_only=True)
        formulas_book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        name = sheet_name or values_book.sheetnames[0]
        if name not in values_book.sheetnames:
            raise TableMappingError(f"Worksheet {name} is missing.")
        values = [list(row) for row in values_book[name].values]
        formulas = [list(row) for row in formulas_book[name].values]
        return values, formulas, name
    except (OSError, ValueError, IndexError, KeyError) as exc:
        if isinstance(exc, TableMappingError):
            raise
        raise TableMappingError("XLSX cannot be read as a workbook.") from exc


def _string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value)


def canonical_table(values: list[list[Any]], formulas: list[list[Any]], *,
                    start_row: int = 1, start_col: int = 1) -> str:
    """One unique line per cell; formulas are retained but do not affect value identity."""
    lines = []
    for row_index, row in enumerate(values, start=start_row):
        formula_row = formulas[row_index - start_row] if row_index - start_row < len(formulas) else []
        for local_index, value in enumerate(row):
            if value is None and (local_index >= len(formula_row) or formula_row[local_index] is None):
                continue
            col = ""
            n = local_index + start_col
            while n:
                n, rem = divmod(n - 1, 26)
                col = chr(65 + rem) + col
            formula = formula_row[local_index] if local_index < len(formula_row) else None
            suffix = f" [formula: {_string(formula)}]" if isinstance(formula, str) and formula.startswith("=") else ""
            lines.append(f"{col}{row_index}: {_string(value)}{suffix}")
    return "\n".join(lines)


def _typed(value: Any, column: TypedColumn) -> dict[str, Any]:
    if value is None or _string(value).strip() == "":
        raise TableMappingError(f"Column {column.header} has an empty value.")
    if column.type == "date":
        if isinstance(value, datetime):
            result = value.date()
        elif isinstance(value, date):
            result = value
        elif isinstance(value, (int, float)):
            result = date(1899, 12, 30) + timedelta(days=int(value))
        else:
            try:
                result = date.fromisoformat(str(value).strip())
            except ValueError as exc:
                raise TableMappingError(f"Column {column.header} needs an ISO date.") from exc
        return {"type": "date", "value": result.isoformat()}
    if column.type == "decimal":
        try:
            number = Decimal(str(value).strip())
        except InvalidOperation as exc:
            raise TableMappingError(f"Column {column.header} needs a decimal number.") from exc
        if not number.is_finite():
            raise TableMappingError(f"Column {column.header} needs a finite decimal number.")
        return {"type": "decimal", "value": str(number), "unit": column.unit}
    result = _string(value).strip()
    if result not in (column.enum_values or []):
        raise TableMappingError(f"Column {column.header} has an unrecognized enum value.")
    return {"type": "enum", "value": result}


def mapped_facts(project_id: str, source: dict[str, Any], mapping: TableMapping,
                 values: list[list[Any]], *, start_row: int = 1, start_col: int = 1) -> list[dict[str, Any]]:
    header_index = mapping.header_row - start_row
    if header_index < 0 or header_index >= len(values):
        raise TableMappingError("Mapped header row is missing. Mapping paused for schema drift.")
    headers = [_string(value).strip() for value in values[header_index]]
    required = [mapping.row_business_key, *(column.header for column in mapping.columns)]
    if len(headers) != len(set(headers)) or any(header not in headers for header in required):
        raise TableMappingError("Mapped header is missing or duplicated. Mapping paused for schema drift.")
    facts = []
    seen_keys = set()
    for row_number, row in enumerate(values[header_index + 1:], start=mapping.header_row + 1):
        if not any(_string(value).strip() for value in row):
            continue
        key_index = headers.index(mapping.row_business_key)
        entity = _string(row[key_index] if key_index < len(row) else None).strip()
        if not entity or entity in seen_keys:
            raise TableMappingError(f"Row {row_number} has a missing or duplicate business key. Mapping paused.")
        seen_keys.add(entity)
        for column in mapping.columns:
            index = headers.index(column.header)
            if index >= len(row):
                raise TableMappingError(f"Row {row_number} lost column {column.header}. Mapping paused for schema drift.")
            typed = _typed(row[index], column)
            letter = ""
            n = index + start_col
            while n:
                n, rem = divmod(n - 1, 26)
                letter = chr(65 + rem) + letter
            quote = f"{letter}{row_number}: {_string(row[index])}"
            snapshot = source["contentSnapshot"]
            # A formula is recorded after the effective value, so the exact
            # cell prefix remains a stable quoted span.
            span = resolve_quote(source["id"], snapshot, quote)
            locator = f"{mapping.sheet_name or 'Table'}!{letter}{row_number}; row key {entity}"
            identifier = sha256(f"{source['id']}:{entity}:{column.predicate}".encode()).hexdigest()[:24]
            fact = Fact.model_validate({
                "id": f"fact_{identifier}", "projectId": project_id,
                "sourceId": source["id"], "entityId": entity,
                "predicate": column.predicate, "value": typed, "modality": "confirmed",
                "evidence": {"sourceId": source["id"], "sourceTitle": source["name"],
                             "nativeUrl": source["nativeUrl"], "quote": quote,
                             "locator": locator, "sourceVersionId": span.version_id,
                             "spanId": span.id},
            })
            facts.append(fact.model_dump(mode="json", by_alias=True))
            if len(facts) > 100:
                raise TableMappingError("Mapping exceeds the 100-fact limit.")
    return facts
